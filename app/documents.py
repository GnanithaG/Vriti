"""Build ATS-friendly Word files (single column, standard headings) and read uploaded resumes."""
import io
import re

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ACCENT = RGBColor(0x2C, 0x47, 0xD6)
MUTED = RGBColor(0x55, 0x5C, 0x6B)


def _base_doc(margin_in: float) -> Document:
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(margin_in))
    sec.top_margin = sec.bottom_margin = Inches(0.5)
    style = d.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10.5)
    style.paragraph_format.space_after = Pt(0)
    return d


def _run(p, text, size=10.5, bold=False, color=None):
    r = p.add_run(str(text or ""))
    r.font.size, r.bold = Pt(size), bold
    if color:
        r.font.color.rgb = color
    return r


def _heading(d, text):
    p = d.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(4)
    _run(p, text.upper(), 10, True, ACCENT)
    # bottom border line
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "2"), ("color", "C9CFDD")):
        bottom.set(qn(f"w:{k}"), v)
    bdr.append(bottom)
    pPr.append(bdr)


def _line(d, left, right, width_in):
    p = d.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(4), Pt(1)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(width_in), WD_TAB_ALIGNMENT.RIGHT)
    _run(p, left, bold=True)
    if right:
        _run(p, "\t" + right, 9.5, color=MUTED)


def _bullet(d, text):
    p = d.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    _run(p, text)


def resume_docx(r: dict, fallback_name: str = "") -> bytes:
    d = _base_doc(0.625)
    width = 8.5 - 2 * 0.625
    p = d.add_paragraph()
    _run(p, r.get("name") or fallback_name, 17, True)
    if r.get("headline"):
        _run(d.add_paragraph(), r["headline"], 10.5, color=RGBColor(0x3A, 0x41, 0x52))
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    _run(p, "  |  ".join(c for c in r.get("contact") or [] if c), 9.5, color=MUTED)
    if r.get("summary"):
        _heading(d, "Summary")
        _run(d.add_paragraph(), r["summary"])
    if r.get("skills"):
        _heading(d, "Skills")
        for g in r["skills"]:
            p = d.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            _run(p, f"{g.get('group', '')}: ", bold=True)
            _run(p, ", ".join(g.get("items") or []))
    if r.get("experience"):
        _heading(d, "Experience")
        for j in r["experience"]:
            loc = f", {j['location']}" if j.get("location") else ""
            _line(d, f"{j.get('title', '')} — {j.get('company', '')}{loc}", j.get("dates", ""), width)
            for b in j.get("bullets") or []:
                _bullet(d, b)
    if r.get("education"):
        _heading(d, "Education")
        for e in r["education"]:
            _line(d, f"{e.get('degree', '')} — {e.get('school', '')}", e.get("dates", ""), width)
            if e.get("details"):
                _run(d.add_paragraph(), e["details"])
    for x in r.get("extra") or []:
        if x.get("items"):
            _heading(d, x.get("heading", ""))
            for i in x["items"]:
                _bullet(d, i)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def letter_docx(text: str) -> bytes:
    d = _base_doc(1.0)
    for para in re.split(r"\n\s*\n", text or ""):
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        _run(p, para.strip(), 11)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def read_resume(data: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs).strip()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(data)).pages).strip()
    return data.decode("utf-8", errors="ignore").strip()
